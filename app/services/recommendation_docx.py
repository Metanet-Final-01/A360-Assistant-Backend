"""추천안(Recommendation)을 사람이 읽는 Word(.docx) 서식 문서로 렌더한다 (FR-17, RPA-296).

기존 `/export`의 JSON은 기계 교환·재적재용이고, 이 모듈은 담당자가 검토·공유·결재에 쓸
**서식 문서**를 만든다. python-docx(업로드 파서가 이미 쓰는 의존성, app/services/parser/docx.py)를
쓰기용으로 재사용 — 새 의존성이 없다.

package/action 표기는 RAG 카탈로그를 따른다(예: `Excel_MS / GoToCell`) — 골드셋 채점·에이전트
계약과 같은 표기라 문서·데이터가 어긋나지 않는다([[a360-agent-handoff]] 정준환 합의).
"""

from __future__ import annotations

import logging
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt

from app.schemas import Recommendation
from app.schemas.recommendation import RecommendedAction

logger = logging.getLogger(__name__)

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# 액션 트리 깊이당 들여쓰기(포인트). 파라미터·근거는 여기에 조금 더 얹는다.
_INDENT_PER_DEPTH = 18
_DETAIL_EXTRA_INDENT = 14


def _pct(v: float | None) -> str:
    return f"{round(v * 100)}%" if v is not None else "—"


def _kv(value) -> str:
    """파라미터 값의 사람용 표기 — None/빈값은 미지정으로."""
    if value is None or value == "":
        return "(미지정)"
    return str(value)


def _add_indented(doc: Document, text: str, indent_pt: float, *, italic: bool = False, bold: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(indent_pt)
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    return p


def _render_actions(doc: Document, actions: list[RecommendedAction], depth: int = 0) -> None:
    """액션 트리를 재귀로 렌더한다 — 컨테이너(Loop/If/Step 등)의 children는 더 깊이 들여쓴다."""
    base = _INDENT_PER_DEPTH * (depth + 1)
    for a in actions:
        head = f"{a.order}. {a.package} / {a.action}"
        if a.label and a.label != a.action:
            head += f" — {a.label}"
        p = _add_indented(doc, head, base, bold=True)
        if a.confidence is not None:
            c = p.add_run(f"   (신뢰도 {_pct(a.confidence)})")
            c.italic = True
        for prm in a.parameters:
            label = prm.label or prm.name
            _add_indented(doc, f"· {label}: {_kv(prm.value)}", base + _DETAIL_EXTRA_INDENT)
        if a.rationale:
            _add_indented(doc, f"근거: {a.rationale}", base + _DETAIL_EXTRA_INDENT, italic=True)
        _render_actions(doc, a.children, depth + 1)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val


def build_recommendation_docx(
    payload,
    *,
    session_id: str,
    version: int,
    source: str | None,
    exported_at: str,
    flow_images: list[bytes] | None = None,
) -> bytes:
    """저장된 추천안 payload(dict 또는 Recommendation)를 .docx 바이트로 렌더한다.

    payload는 저장 시 Recommendation으로 검증된 것이지만, 여기서도 model_validate로 타입을
    확보한다(구조가 깨졌으면 ValidationError로 드러난다 — 호출부가 처리).

    flow_images: 프론트가 캡처한 흐름도 PNG/JPEG 바이트 목록(선택, RPA-296·RPA-334). 있으면
    "추천 흐름" 머리 아래에 순서대로 임베드하되 장 **사이**에 페이지 나눔을 넣어 각 장을 새
    페이지에 싣는다(긴 흐름도를 페이지 단위로 나눠 캡처해 오기 때문). 흐름도는 프론트(FR-18)가
    트리에서 그리므로 백엔드가 서버에서 캡처할 수 없어, 호출부가 이미지 목록을 넘겨준다.
    0장이면 이미지 없이 데이터 문서로 동작한다(하위호환).
    """
    rec = payload if isinstance(payload, Recommendation) else Recommendation.model_validate(payload)

    doc = Document()
    doc.add_heading("A360 작업 추천안", level=0)

    meta = doc.add_paragraph()
    meta.add_run(f"세션 {session_id} · 버전 v{version}").bold = True
    bits = [f"편집 출처: {source or '—'}", f"내보낸 시각: {exported_at}"]
    if rec.flow_confidence is not None:
        bits.append(f"흐름도 신뢰도: {_pct(rec.flow_confidence)}")
    doc.add_paragraph(" · ".join(bits))

    # 개요 — spec.goal 우선, 없으면 notes
    goal = (rec.spec.goal if rec.spec and rec.spec.goal else "") or (rec.notes or "")
    if goal:
        doc.add_heading("개요", level=1)
        doc.add_paragraph(goal)

    # 요구사항 (FlowSpec)
    if rec.spec and rec.spec.requirements:
        doc.add_heading("요구사항", level=1)
        _add_table(
            doc,
            ["ID", "요구사항", "우선순위", "출처"],
            [[r.req_id, r.text, r.priority, r.source] for r in rec.spec.requirements],
        )

    # 입력 / 출력
    if rec.spec and (rec.spec.inputs or rec.spec.outputs):
        doc.add_heading("입력 · 출력", level=1)
        if rec.spec.inputs:
            doc.add_paragraph("입력: " + ", ".join(rec.spec.inputs))
        if rec.spec.outputs:
            doc.add_paragraph("산출물: " + ", ".join(rec.spec.outputs))

    # 실행 시점 (trigger)
    if rec.trigger:
        doc.add_heading("실행 시점", level=1)
        t = rec.trigger
        head = t.title + (f"  ({t.package})" if t.package else "")
        doc.add_paragraph().add_run(head).bold = True  # add_run은 빈 문자열에도 안전(runs[0] 인덱싱 회피)
        if t.reason:
            doc.add_paragraph(f"근거: {t.reason}")
        if t.setup_hint:
            doc.add_paragraph(f"설정: {t.setup_hint}")

    # 변수
    if rec.variables:
        doc.add_heading("변수", level=1)
        _add_table(
            doc,
            ["이름", "타입", "방향", "설명"],
            [[v.name, v.type, v.direction, v.description or ""] for v in rec.variables],
        )

    # 추천 흐름 — 본문
    doc.add_heading("추천 흐름", level=1)
    # 프론트가 캡처한 UI 흐름도를 시각 요약으로 먼저 싣고, 아래에 단계별 텍스트를 잇는다.
    # 긴 흐름도는 페이지 높이 단위로 여러 장 캡처해 오므로 장 **사이**에만 페이지 나눔을 넣어
    # 각 장을 새 페이지에서 시작하게 한다 — Word는 페이지보다 큰 인라인 그림을 이어 그리지 않고
    # 경계에서 잘라버린다(RPA-334). 첫 장 앞·마지막 장 뒤엔 넣지 않는다(뒤 텍스트와 붙어도 무방).
    total = len(flow_images) if flow_images else 0
    for idx, image in enumerate(flow_images or []):
        if idx > 0:
            doc.add_page_break()
        try:
            doc.add_picture(BytesIO(image), width=Inches(6.3))
            caption = "흐름도 (편집 화면 기준)" if total == 1 else f"흐름도 ({idx + 1}/{total})"
            doc.add_paragraph().add_run(caption).italic = True
        except Exception:  # noqa: BLE001 — 한 장이 깨져도 문서 생성·다른 장은 죽이지 않는다
            logger.warning("흐름도 이미지 임베드 실패 — 문구로 대체 (%d/%d)", idx + 1, total, exc_info=True)
            doc.add_paragraph(f"(흐름도 이미지를 표시할 수 없습니다. {idx + 1}/{total})")
    if not rec.steps:
        doc.add_paragraph("(흐름 단계가 없습니다.)")
    for i, step in enumerate(rec.steps, start=1):
        title = step.label or step.step_id
        doc.add_heading(f"{i}. {title}", level=2)
        if step.description:
            doc.add_paragraph(step.description)
        _render_actions(doc, step.actions)

    # 확인 필요 — 질문 카드
    if rec.needs_input:
        doc.add_heading("확인 필요", level=1)
        for card in rec.needs_input:
            mark = "[필수] " if card.blocking else ""
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{mark}{card.question}").bold = True
            if card.why:
                _add_indented(doc, f"이유: {card.why}", _INDENT_PER_DEPTH, italic=True)
            if card.default is not None:
                _add_indented(doc, f"시안값: {_kv(card.default)}", _INDENT_PER_DEPTH)

    # 전제 (assumptions) — 생성이 임의로 정한 전제는 명시적으로 드러낸다
    if rec.spec and rec.spec.assumptions:
        doc.add_heading("전제", level=1)
        for a in rec.spec.assumptions:
            doc.add_paragraph(a, style="List Bullet")

    # 주의사항 (notes) — 개요로 이미 쓰지 않았을 때만
    if rec.notes and rec.notes != goal:
        doc.add_heading("주의사항", level=1)
        doc.add_paragraph(rec.notes)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
