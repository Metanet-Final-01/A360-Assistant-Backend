"""추천안 내보내기 엔드포인트 테스트 (RPA-79, FR-17)."""

import uuid
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

import app.api.sessions as sessions_api
from app.db import get_db
from app.main import app

SID = uuid.uuid4()


def _rec() -> dict:
    return {
        "schema_version": "1.0",
        "steps": [{"step_id": "step-1", "actions": [
            {"order": 1, "package": "Browser", "action": "openbrowser", "label": "열기",
             "parameters": [], "children": []}]}],
        "variables": [], "notes": "",
    }


class FakeDB:
    def __init__(self, session=None, row=None):
        self.session = session
        self.row = row

    def get(self, model, key):
        return self.session

    def execute(self, stmt):
        return SimpleNamespace(scalar_one_or_none=lambda: self.row)


@pytest.fixture(autouse=True)
def _cleanup():
    yield
    app.dependency_overrides.clear()


def _override(db, user=None):
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[sessions_api.get_optional_user] = lambda: user


def test_export_returns_download_envelope():
    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=3, source="drag", payload=_rec())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/3/export")

    assert r.status_code == 200
    cd = r.headers["content-disposition"]
    assert "attachment" in cd and "v3.json" in cd  # 다운로드 헤더
    body = r.json()
    assert body["recommendation_version"] == 3 and body["source"] == "drag"
    assert body["recommendation"]["steps"][0]["step_id"] == "step-1"  # 트리 그대로
    assert "exported_at" in body


def test_export_404_missing_version():
    session = SimpleNamespace(id=SID, user_id=None)
    _override(FakeDB(session=session, row=None))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/99/export")
    assert r.status_code == 404


def test_export_blocks_non_owner():
    session = SimpleNamespace(id=SID, user_id=uuid.uuid4())
    _override(FakeDB(session=session, row=None), user=SimpleNamespace(id=uuid.uuid4()))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/1/export")
    assert r.status_code == 403


def test_export_non_int_version_422():
    """version은 int — /recommendations/latest 등과 라우트가 안 섞인다."""
    session = SimpleNamespace(id=SID, user_id=None)
    _override(FakeDB(session=session, row=None))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/notanint/export")
    assert r.status_code == 422


# --- docx 서식 문서 내보내기 (RPA-296, FR-17) ---

def _rec_rich() -> dict:
    """docx 렌더의 주요 섹션(개요·요구사항·입출력·trigger·변수·중첩 액션·질문카드·전제)을 모두 밟는다."""
    return {
        "schema_version": "1.0",
        "flow_confidence": 0.82,
        "steps": [{
            "step_id": "step-1", "label": "엑셀 열기", "description": "대상 파일을 연다",
            "actions": [{
                "order": 1, "package": "Excel_MS", "action": "OpenWorkbook", "label": "통합문서 열기",
                "parameters": [{"name": "path", "label": "파일 경로", "value": "C:/data.xlsx"}],
                "rationale": "업무정의서의 '엑셀 파일' 근거", "confidence": 0.9,
                "children": [{
                    "order": 2, "package": "Excel_MS", "action": "GoToCell", "label": "셀 이동",
                    "parameters": [{"name": "cellOption", "value": "A1"}], "children": [],
                }],
            }],
        }],
        "variables": [{"name": "vFile", "type": "STRING", "direction": "input", "description": "입력 파일"}],
        "trigger": {"kind": "schedule", "title": "매일 오전 9시", "reason": "'매일 아침' 표현",
                    "setup_hint": "Control Room 예약"},
        "needs_input": [{"card_id": "c1", "kind": "missing_param", "question": "파일 경로를 알려주세요",
                         "why": "열 파일이 필요", "blocking": True, "default": None}],
        "spec": {"goal": "매일 엑셀 집계 자동화",
                 "requirements": [{"req_id": "req-1", "text": "엑셀 열기", "priority": "must", "source": "doc"}],
                 "inputs": ["엑셀 파일"], "outputs": ["집계 결과"],
                 "assumptions": ["Knox 메일은 Email 패키지 기준"]},
        "notes": "테스트 주의사항",
    }


def test_export_docx_returns_formatted_document():
    """format=docx → 실제로 열리는 .docx이고, 카탈로그 표기·중첩 액션·주요 섹션이 담긴다."""
    from io import BytesIO

    from docx import Document

    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=5, source="agent", payload=_rec_rich())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/5/export", params={"format": "docx"})

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    cd = r.headers["content-disposition"]
    assert "attachment" in cd and "v5.docx" in cd
    assert len(r.content) > 0

    doc = Document(BytesIO(r.content))
    blob = "\n".join(p.text for p in doc.paragraphs)
    blob += "\n" + "\n".join(cell.text for t in doc.tables for row_ in t.rows for cell in row_.cells)
    assert "A360 작업 추천안" in blob
    assert "Excel_MS / OpenWorkbook" in blob     # 카탈로그 표기(package/action)
    assert "Excel_MS / GoToCell" in blob          # 컨테이너 children도 재귀 렌더
    assert "매일 오전 9시" in blob                 # trigger
    assert "vFile" in blob                         # 변수표
    assert "매일 엑셀 집계 자동화" in blob          # 개요(spec.goal)
    assert "파일 경로를 알려주세요" in blob          # 질문카드
    assert "Knox 메일은 Email 패키지 기준" in blob   # 전제(assumptions)


def test_export_json_is_default_and_unchanged():
    """format 미지정·명시 json 모두 기존 JSON 봉투 — 하위호환."""
    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=3, source="drag", payload=_rec())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/3/export", params={"format": "json"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/json")
    assert r.json()["recommendation"]["steps"][0]["step_id"] == "step-1"


def test_export_invalid_format_422():
    """format은 json|docx만 — 그 외는 422(pdf 등 미지원 포맷 오인 방지)."""
    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=1, source="drag", payload=_rec())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/1/export", params={"format": "pdf"})
    assert r.status_code == 422


def test_export_docx_render_failure_returns_standard_500():
    """렌더 실패(잘못된 payload)는 트레이스백 500이 아니라 표준 DOCX_RENDER_FAILED (Qodo #421, GET 경로)."""
    session = SimpleNamespace(id=SID, user_id=None)
    # steps 없음 → Recommendation 검증 실패 → build가 예외 → 엔드포인트가 표준 에러로 저하
    row = SimpleNamespace(id=uuid.uuid4(), version=1, source="drag", payload={"schema_version": "1.0"})
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.get(f"/api/sessions/{SID}/recommendations/1/export", params={"format": "docx"})
    assert r.status_code == 500
    assert r.json()["detail"]["code"] == "DOCX_RENDER_FAILED"


# --- 프론트 캡처 흐름도 임베드 (RPA-296, POST export/docx) ---

def _png_bytes() -> bytes:
    from io import BytesIO

    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (160, 90), (210, 225, 255)).save(buf, "PNG")
    return buf.getvalue()


def _page_break_count(doc) -> int:
    """문서의 명시적 페이지 나눔(<w:br w:type="page"/>) 개수 — 장 사이 나눔을 검증한다."""
    from docx.oxml.ns import qn

    return sum(
        1 for br in doc.element.body.iter(qn("w:br"))
        if br.get(qn("w:type")) == "page"
    )


def test_export_docx_post_embeds_multiple_flow_images():
    """POST + 캡처 PNG 여러 장 → 각 장이 임베드되고 장 **사이**에만 페이지 나눔이 들어간다 (RPA-334)."""
    from io import BytesIO

    from docx import Document

    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=5, source="agent", payload=_rec_rich())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.post(
            f"/api/sessions/{SID}/recommendations/5/export/docx",
            files=[
                ("flow_images", ("flow1.png", _png_bytes(), "image/png")),
                ("flow_images", ("flow2.png", _png_bytes(), "image/png")),
                ("flow_images", ("flow3.png", _png_bytes(), "image/png")),
            ],
        )
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "v5.docx" in r.headers["content-disposition"]
    doc = Document(BytesIO(r.content))
    assert len(doc.inline_shapes) == 3            # 세 장 모두 임베드
    assert _page_break_count(doc) == 2            # 장 사이에만(첫 장 앞·마지막 장 뒤엔 없음)
    blob = "\n".join(p.text for p in doc.paragraphs)
    assert "흐름도 (1/3)" in blob and "흐름도 (3/3)" in blob  # 장 번호 캡션
    assert "Excel_MS / OpenWorkbook" in blob      # 데이터 섹션도 함께 렌더


def test_export_docx_post_single_image_no_page_break():
    """1장이면 페이지 나눔 없이 기존과 동일하게 임베드하고 캡션은 번호 없음(하위호환) (RPA-334)."""
    from io import BytesIO

    from docx import Document

    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=5, source="agent", payload=_rec_rich())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.post(
            f"/api/sessions/{SID}/recommendations/5/export/docx",
            files=[("flow_images", ("flow.png", _png_bytes(), "image/png"))],
        )
    assert r.status_code == 200
    doc = Document(BytesIO(r.content))
    assert len(doc.inline_shapes) == 1
    assert _page_break_count(doc) == 0
    blob = "\n".join(p.text for p in doc.paragraphs)
    assert "흐름도 (편집 화면 기준)" in blob  # 단일 장 캡션은 번호 없음


def test_export_docx_post_without_image_is_data_doc():
    """이미지 없이 POST해도 데이터 문서로 동작한다(이미지 임베드 없음)."""
    from io import BytesIO

    from docx import Document

    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=5, source="agent", payload=_rec())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.post(f"/api/sessions/{SID}/recommendations/5/export/docx")
    assert r.status_code == 200
    doc = Document(BytesIO(r.content))
    assert len(doc.inline_shapes) == 0


def test_docx_trigger_empty_title_does_not_crash():
    """trigger.title이 빈 문자열이어도 렌더가 크래시하지 않는다 (Qodo #421 — runs[0] 인덱싱 회피)."""
    from app.services.recommendation_docx import build_recommendation_docx

    payload = {"schema_version": "1.0", "steps": [],
               "trigger": {"kind": "trigger", "title": ""}}
    content = build_recommendation_docx(payload, session_id="s", version=1, source=None, exported_at="t")
    assert content and len(content) > 0


def test_export_docx_post_rejects_non_image():
    """여러 장 중 하나라도 PNG/JPEG 매직바이트가 아니면 400 — 장마다 개별 검증 (RPA-334)."""
    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=1, source="drag", payload=_rec())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.post(
            f"/api/sessions/{SID}/recommendations/1/export/docx",
            files=[
                ("flow_images", ("ok.png", _png_bytes(), "image/png")),
                ("flow_images", ("evil.txt", b"not really an image", "image/png")),
            ],
        )
    assert r.status_code == 400
    assert r.json()["detail"]["code"] == "INVALID_IMAGE"


def test_export_docx_post_rejects_too_many_images():
    """이미지 개수가 상한(20장)을 넘으면 413 TOO_MANY_IMAGES (RPA-334)."""
    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=1, source="drag", payload=_rec())
    _override(FakeDB(session=session, row=row))
    files = [("flow_images", (f"f{i}.png", _png_bytes(), "image/png")) for i in range(21)]
    with TestClient(app) as c:
        r = c.post(f"/api/sessions/{SID}/recommendations/1/export/docx", files=files)
    assert r.status_code == 413
    assert r.json()["detail"]["code"] == "TOO_MANY_IMAGES"


def test_export_docx_post_rejects_oversized_total(monkeypatch):
    """장별 8MB 이하라도 합계가 상한을 넘으면 413 IMAGES_TOO_LARGE (RPA-334, Qodo #445 메모리 폭증)."""
    # 합계 상한을 한 장 크기로 낮춰 2장이면 합계 초과하도록(테스트에 48MB를 만들지 않기 위함).
    monkeypatch.setattr(sessions_api, "_MAX_FLOW_IMAGES_TOTAL_BYTES", len(_png_bytes()))
    session = SimpleNamespace(id=SID, user_id=None)
    row = SimpleNamespace(id=uuid.uuid4(), version=1, source="drag", payload=_rec())
    _override(FakeDB(session=session, row=row))
    with TestClient(app) as c:
        r = c.post(
            f"/api/sessions/{SID}/recommendations/1/export/docx",
            files=[
                ("flow_images", ("a.png", _png_bytes(), "image/png")),
                ("flow_images", ("b.png", _png_bytes(), "image/png")),
            ],
        )
    assert r.status_code == 413
    assert r.json()["detail"]["code"] == "IMAGES_TOO_LARGE"
